"""Declare, export, and read workbooks for a person's file step.

Columns are fixed (export only), text (stripped on import), or list (closed
values with a dropdown). Lists normalize by stripping and case-folding before
the optional caller normalizer; blank is accepted only with allow_blank.
A list's baseline names a fixed column: changed cells become bold red. Tab
fills map list values to colours; start names a list column and its continuation
value, with every other nonblank value filling the row blue. A wrapped column
splits long text into rows whose id and editable cells remain empty.

Import never writes. Imported rows are keyed by worksheet row number and hold
the id and editable columns only. Findings hold coordinates and kinds alone:
not-in-list, continuation-edited, column-missing, tab-missing, and
index-missing. A missing heading uses the header row; a missing tab or index
uses row 0 and the id column's name. Explicit tab declarations select which tabs
to read; without them, index ids select tabs for existence checks.
"""

import math
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_ROW = 4
CHUNK = 2500
TEXT_WIDTH = 110
START_FILL = "DDEBF7"
UNKNOWN_FILL = "FCE4D6"
BOLD = Font(bold=True)
LINK = Font(color="0563C1", underline="single")
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
CHANGED_FONT = Font(bold=True, color="C00000")

NOT_IN_LIST = "not-in-list"
CONTINUATION_EDITED = "continuation-edited"
TAB_MISSING = "tab-missing"
INDEX_MISSING = "index-missing"
COLUMN_MISSING = "column-missing"


@dataclass(frozen=True)
class Column:
    """One table column and its export and import rules."""

    name: str
    kind: Literal["fixed", "text", "list"] = "fixed"
    values: Sequence[str] = ()
    baseline: str | None = None
    width: float | None = None
    wrap: bool = False
    number_format: str | None = None
    allow_blank: bool = False
    normalizer: Callable[[str], str] | None = None
    font_color: str | None = None
    error_title: str | None = None
    error: str | None = None


@dataclass(frozen=True)
class Cell:
    """A preamble cell, optionally bold or linked."""

    row: int
    column: int
    value: object
    hyperlink: str | None = None
    bold: bool = False


@dataclass(frozen=True)
class Tab:
    """One unit's preamble, table, and conditional row colours."""

    title: str
    columns: Sequence[Column]
    id_column: str
    preamble: Sequence[Cell] = ()
    rows: Sequence[Sequence[object]] = ()
    fills: Mapping[str, str] = field(default_factory=dict)
    fill_column: str | None = None
    start: tuple[str, str] | None = None
    start_fill: str = START_FILL


@dataclass(frozen=True)
class Index:
    """The index table, with unit links and optional external link labels."""

    title: str
    columns: Sequence[Column]
    id_column: str
    rows: Sequence[Sequence[object]] = ()
    hyperlinks: Mapping[str, str] = field(default_factory=dict)


@dataclass(frozen=True)
class Book:
    """A guide, index, and unit tabs, or the same shape without data rows."""

    guide_title: str
    guide_lines: Sequence[str]
    index: Index
    tabs: Sequence[Tab] = ()
    header_row: int = HEADER_ROW
    chunk_size: int = CHUNK
    guide_headings: Sequence[str] = ()


@dataclass(frozen=True)
class Finding:
    """A content-free cell location and rule to correct."""

    tab: str
    row: int
    column: str
    kind: str


@dataclass(frozen=True)
class Imported:
    """Editable values by worksheet row, and all validation findings."""

    index_rows: Mapping[int, Mapping[str, object]] = field(repr=False)
    tab_rows: Mapping[str, Mapping[int, Mapping[str, object]]] = field(repr=False)
    findings: tuple[Finding, ...]


class WorkbookError(Exception):
    """An export refusal with its correction."""

    def __init__(self, problem: str, fix: str) -> None:
        super().__init__(problem)
        self.problem = problem
        self.fix = fix


def _put(cell: openpyxl.cell.cell.Cell, value: object) -> None:
    if isinstance(value, str):
        cell.value = ILLEGAL_CHARACTERS_RE.sub(" ", value)
        cell.data_type = "s"
    else:
        cell.value = value


def _headings(tab: openpyxl.worksheet.worksheet.Worksheet,
              columns: Sequence[Column], row: int) -> None:
    for position, column in enumerate(columns, 1):
        cell = tab.cell(row, position)
        _put(cell, column.name)
        cell.font = BOLD
        width = column.width if column.width is not None else TEXT_WIDTH if column.wrap else None
        if width is not None:
            tab.column_dimensions[get_column_letter(position)].width = width


def _validations(tab: openpyxl.worksheet.worksheet.Worksheet,
                 columns: Sequence[Column]) -> dict[int, openpyxl.worksheet.datavalidation.DataValidation]:
    validations = {}
    for position, column in enumerate(columns, 1):
        if column.kind != "list":
            continue
        validation = openpyxl.worksheet.datavalidation.DataValidation(
            type="list", formula1='"' + ",".join(column.values) + '"',
            allow_blank=column.allow_blank,
        )
        if column.error_title is not None:
            validation.showErrorMessage = True
            validation.errorTitle = column.error_title
            validation.error = column.error
        tab.add_data_validation(validation)
        validations[position] = validation
    return validations


def _chunks(text: str, size: int) -> list[str]:
    parts = []
    while len(text) > size:
        cut = text.rfind(" ", 0, size)
        cut = cut if cut > size // 2 else size
        parts.append(text[:cut])
        text = text[cut:].lstrip()
    parts.append(text)
    return parts


def _conditions(tab: openpyxl.worksheet.worksheet.Worksheet, shape: Tab,
                first: int, last: int) -> None:
    if last < first:
        return
    positions = {column.name: get_column_letter(i) for i, column in enumerate(shape.columns, 1)}
    span = f"A{first}:{get_column_letter(len(shape.columns))}{last}"
    fill_column = shape.fill_column or (shape.start[0] if shape.start else next(
        (column.name for column in shape.columns if column.kind == "list"), None,
    ))
    for value, colour in shape.fills.items():
        assert fill_column is not None
        escaped = value.replace('"', '""')
        tab.conditional_formatting.add(span, FormulaRule(
            formula=[f'${positions[fill_column]}{first}="{escaped}"'],
            fill=PatternFill("solid", start_color=colour),
        ))
    if shape.start:
        name, continuation = shape.start
        cell = f"${positions[name]}{first}"
        escaped = continuation.replace('"', '""')
        tab.conditional_formatting.add(span, FormulaRule(
            formula=[f'AND({cell}<>"",{cell}<>"{escaped}")'],
            fill=PatternFill("solid", start_color=shape.start_fill),
        ))
    for column in shape.columns:
        if column.kind == "list" and column.baseline:
            letter = positions[column.name]
            tab.conditional_formatting.add(f"{letter}{first}:{letter}{last}", FormulaRule(
                formula=[f'AND(${positions[shape.id_column]}{first}<>"",'
                         f'${letter}{first}<>${positions[column.baseline]}{first})'],
                font=CHANGED_FONT,
            ))


def export(book: Book, out: Path, *, checkout: Path) -> tuple[int, int]:
    """Write an absent file outside checkout; return unit-tab and logical-row counts."""
    if out.exists() or out.is_symlink():
        raise WorkbookError("output exists", "Choose another name so a checked workbook is never overwritten.")
    if out.resolve().is_relative_to(checkout.resolve()):
        raise WorkbookError("output is inside the checkout", "Choose a path outside the checkout.")

    written = openpyxl.Workbook()
    try:
        guide = written.active
        guide.title = book.guide_title
        guide.column_dimensions["A"].width = 130
        for row, text in enumerate(book.guide_lines, 1):
            cell = guide.cell(row, 1)
            _put(cell, text)
            cell.font = BOLD if text in book.guide_headings else Font()

        index = written.create_sheet(book.index.title)
        _headings(index, book.index.columns, 1)
        index.freeze_panes = "A2"
        validations = _validations(index, book.index.columns)
        for row, values in enumerate(book.index.rows, 2):
            for position, (column, value) in enumerate(zip(book.index.columns, values, strict=True), 1):
                cell = index.cell(row, position)
                _put(cell, value)
                if column.name == book.index.id_column:
                    title = str(value).replace("'", "''")
                    cell.hyperlink = f"#'{title}'!A1"
                    cell.font = LINK
                elif column.name in book.index.hyperlinks and value:
                    _put(cell, book.index.hyperlinks[column.name])
                    cell.hyperlink = str(value)
                    cell.font = LINK
                if position in validations:
                    validations[position].add(cell)
        index.auto_filter.ref = f"A1:{get_column_letter(len(book.index.columns))}{len(book.index.rows) + 1}"

        for shape in book.tabs:
            tab = written.create_sheet(shape.title)
            for item in shape.preamble:
                cell = tab.cell(item.row, item.column)
                _put(cell, item.value)
                if item.hyperlink:
                    cell.hyperlink = item.hyperlink
                    cell.font = LINK
                elif item.bold:
                    cell.font = BOLD
            _headings(tab, shape.columns, book.header_row)
            tab.freeze_panes = tab.cell(book.header_row + 1, 2)
            validations = _validations(tab, shape.columns)
            row = book.header_row + 1
            for values in shape.rows:
                height = 1
                for position, (column, value) in enumerate(zip(shape.columns, values, strict=True), 1):
                    cell = tab.cell(row, position)
                    if column.wrap:
                        parts = _chunks("" if value is None else str(value), book.chunk_size)
                        height = max(height, len(parts))
                        width = column.width if column.width is not None else TEXT_WIDTH
                        for offset, part in enumerate(parts):
                            wrapped = tab.cell(row + offset, position)
                            _put(wrapped, part if offset == 0 else "… " + part)
                            wrapped.alignment = WRAP
                            lines = max(1, math.ceil(len(part) / (width * 0.95)))
                            tab.row_dimensions[row + offset].height = min(409.0, 15.0 * lines + 2)
                    else:
                        _put(cell, value)
                        cell.alignment = TOP
                    if column.number_format:
                        cell.number_format = column.number_format
                    if column.font_color:
                        cell.font = Font(color=column.font_color)
                    if position in validations:
                        validations[position].add(cell)
                row += height
            _conditions(tab, shape, book.header_row + 1, row - 1)
        written.active = 0
        with out.open("xb") as stream:
            try:
                written.save(stream)
            except BaseException:
                out.unlink()
                raise
    finally:
        written.close()
    return len(book.tabs), sum(len(tab.rows) for tab in book.tabs)


def _read_rows(tab: openpyxl.worksheet.worksheet.Worksheet, columns: Sequence[Column],
               id_column: str, header: int, findings: list[Finding]) -> dict[int, Mapping[str, object]]:
    positions = {cell.value: cell.column for cell in tab[header]}
    rows: dict[int, Mapping[str, object]] = {}
    missing = [column.name for column in columns
               if column.name not in positions and (column.kind != "fixed" or column.name == id_column)]
    if missing:
        findings.extend(Finding(tab.title, header, name, COLUMN_MISSING) for name in missing)
        return rows
    for row in range(header + 1, tab.max_row + 1):
        identity = tab.cell(row, positions[id_column]).value
        continuation = identity is None or str(identity).strip() == ""
        values: dict[str, object] = {id_column: identity}
        for column in columns:
            if column.kind == "fixed":
                continue
            value = tab.cell(row, positions[column.name]).value
            if continuation:
                if value is not None and value != "":
                    findings.append(Finding(tab.title, row, column.name, CONTINUATION_EDITED))
                continue
            text = "" if value is None else str(value).strip()
            if column.kind == "list":
                text = text.casefold()
                if column.normalizer:
                    text = column.normalizer(text)
                if text not in column.values and not (column.allow_blank and text == ""):
                    findings.append(Finding(tab.title, row, column.name, NOT_IN_LIST))
            values[column.name] = text
        if not continuation:
            rows[row] = values
    return rows


def import_(path: Path, shape: Book) -> Imported:
    """Read declared editable cells and collect every located, content-free finding."""
    findings: list[Finding] = []
    read = openpyxl.load_workbook(path, data_only=True)
    try:
        if shape.index.title not in read.sheetnames:
            return Imported({}, {}, (Finding(shape.index.title, 0, shape.index.id_column, INDEX_MISSING),))
        index_rows = _read_rows(read[shape.index.title], shape.index.columns, shape.index.id_column, 1, findings)
        tab_rows: dict[str, Mapping[int, Mapping[str, object]]] = {}
        tabs = shape.tabs or tuple(Tab(str(row[shape.index.id_column]), (), shape.index.id_column)
                                   for row in index_rows.values())
        for tab in tabs:
            if tab.title not in read.sheetnames:
                findings.append(Finding(tab.title, 0, tab.id_column, TAB_MISSING))
            elif tab.columns:
                tab_rows[tab.title] = _read_rows(read[tab.title], tab.columns, tab.id_column, shape.header_row, findings)
            else:
                tab_rows[tab.title] = {}
        return Imported(index_rows, tab_rows, tuple(findings))
    finally:
        read.close()
